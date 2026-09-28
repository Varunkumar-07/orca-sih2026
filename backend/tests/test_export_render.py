"""Pytest coverage for backend/services/export_render.py (previously
untested — added as part of the codebase audit's P1 coverage gaps).

No network/Copernicus involved at all — this module only renders data
get_historical_analytics already produced, so tests build that dict shape
directly rather than going through the live fetch pipeline.

Covers:
  - _build_rows: (label, date, value) triples, grouped/sorted by variable
    then chronological; moon_phase shows its display name instead of the
    raw numeric fraction; a None value renders as an em dash, not "None"
  - render_pdf: produces real, non-trivial PDF bytes (%PDF magic header)
    without raising, for both a normal series and an empty one
  - render_docx: produces a real, re-openable .docx (python-docx can read
    it back) whose paragraphs actually contain the zone id, coordinates,
    and date range passed in — not just "doesn't crash"

Run from project root:  PYTHONPATH=. pytest   or   pytest
"""
from __future__ import annotations

import io

from docx import Document

from backend.services import export_render as svc


def _data(series: dict | None = None) -> dict:
    return {
        "lat": 13.08,
        "lon": 80.27,
        "start_date": "2026-01-01",
        "end_date": "2026-01-02",
        "series": series
        if series is not None
        else {
            "sst_celsius": {"unit": "°C", "dates": ["2026-01-01", "2026-01-02"], "values": [28.5, None], "status": "partial"},
            "wind_kmh": {"unit": "km/h", "dates": ["2026-01-01"], "values": [12.0], "status": "ok"},
        },
    }


# ---------------------------------------------------------------------------
# _build_rows
# ---------------------------------------------------------------------------


def test_build_rows_groups_by_variable_then_chronological():
    rows = svc._build_rows(_data())

    # Alphabetical by key ("sst_celsius" before "wind_kmh"), chronological within each.
    assert rows == [
        ("Sea Surface Temperature", "2026-01-01", "28.5 °C"),
        ("Sea Surface Temperature", "2026-01-02", "—"),  # None value -> em dash, not "None °C"
        ("Wind Speed", "2026-01-01", "12.0 km/h"),
    ]


def test_build_rows_moon_phase_shows_display_name_not_raw_fraction():
    series = {
        "moon_phase": {
            "unit": "",
            "dates": ["2026-01-01"],
            "values": [0.02],
            "status": "ok",
            "names": ["New Moon"],
        }
    }
    rows = svc._build_rows(_data(series))
    assert rows == [("Moon Phase", "2026-01-01", "New Moon")]


# ---------------------------------------------------------------------------
# render_pdf
# ---------------------------------------------------------------------------


def test_render_pdf_produces_valid_pdf_bytes():
    pdf_bytes = svc.render_pdf("KOCHI-PFZ-001", _data())
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 500  # a real rendered document, not an empty shell


def test_render_pdf_handles_empty_series_without_raising():
    pdf_bytes = svc.render_pdf("EMPTY-ZONE", _data(series={}))
    assert pdf_bytes.startswith(b"%PDF")


# ---------------------------------------------------------------------------
# render_docx
# ---------------------------------------------------------------------------


def test_render_docx_produces_valid_reopenable_docx_with_real_content():
    docx_bytes = svc.render_docx("KOCHI-PFZ-001", _data())
    assert docx_bytes[:2] == b"PK"  # docx is a zip container

    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "KOCHI-PFZ-001" in full_text
    assert "13.0800, 80.2700" in full_text
    assert "2026-01-01 to 2026-01-02" in full_text

    # One data table with a 3-column header + one row per (variable, date) pair.
    assert len(doc.tables) == 1
    header = [c.text for c in doc.tables[0].rows[0].cells]
    assert header == ["Variable", "Date", "Value"]
    assert len(doc.tables[0].rows) == 1 + 3  # header + 2 sst rows + 1 wind row


def test_render_docx_handles_empty_series_without_raising():
    docx_bytes = svc.render_docx("EMPTY-ZONE", _data(series={}))
    assert docx_bytes[:2] == b"PK"
    doc = Document(io.BytesIO(docx_bytes))
    assert len(doc.tables[0].rows) == 1  # header only, no data rows
